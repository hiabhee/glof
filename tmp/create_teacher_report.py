from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path('docs/reports/glacierlens-south-lhonak-pilot-report-2026-09-27.docx')
BLUE = '17365D'
PALE_BLUE = 'EAF2F8'
PALE_GREY = 'F5F7F9'
GRID = 'D9D9D9'


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders_el = tc_pr.first_child_found_in('w:tcBorders')
    if borders_el is None:
        borders_el = OxmlElement('w:tcBorders')
        tc_pr.append(borders_el)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        element = borders_el.find(qn(f'w:{edge}'))
        if element is None:
            element = OxmlElement(f'w:{edge}')
            borders_el.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '4')
        element.set(qn('w:color'), GRID)


def set_cell_text(cell, text, bold=False, color='000000', size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = 'Aptos'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    borders(cell)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.autofit = False
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell, BLUE)
        set_cell_text(cell, header, bold=True, color='FFFFFF', size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        if widths:
            cell.width = Inches(widths[i])
    for row_i, row in enumerate(rows):
        cells = table.add_row().cells
        for col_i, value in enumerate(row):
            if row_i % 2 == 1:
                shade(cells[col_i], PALE_GREY)
            set_cell_text(cells[col_i], value, size=9)
            if widths:
                cells[col_i].width = Inches(widths[col_i])
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def paragraph(doc, text='', bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.12
    if bold_lead and text.startswith(bold_lead):
        first = p.add_run(bold_lead)
        first.bold = True
        first.font.size = Pt(11)
        rest = p.add_run(text[len(bold_lead):])
        rest.font.size = Pt(11)
    else:
        p.add_run(text).font.size = Pt(11)
    for run in p.runs:
        run.font.name = 'Aptos'
        run._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
        run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.name = 'Aptos Display' if level == 1 else 'Aptos'
    return p


def bullet(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.26)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(4)
    marker = p.add_run('• ')
    marker.font.size = Pt(10.5)
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    r.font.name = 'Aptos'
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
styles['Normal'].font.size = Pt(11)
for style_name in ('Title', 'Heading 1', 'Heading 2'):
    styles[style_name].font.color.rgb = RGBColor(0, 0, 0)

# Word's built-in Title style can carry a blue bottom border. Remove it so the
# title remains plain black text as required for this report.
title_style_ppr = styles['Title']._element.get_or_add_pPr()
title_border = title_style_ppr.find(qn('w:pBdr'))
if title_border is not None:
    title_style_ppr.remove(title_border)

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title_run = title.add_run('GlacierLens South Lhonak Glacier Pilot Project Report')
title_run.font.color.rgb = RGBColor(0, 0, 0)
title_run.font.name = 'Aptos Display'

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(16)
subtitle_run = subtitle.add_run('Experimental geospatial and machine learning pilot for teacher review\n27 September 2026')
subtitle_run.font.size = Pt(12)
subtitle_run.font.color.rgb = RGBColor(80, 80, 80)

heading(doc, 'Project summary')
paragraph(doc, 'GlacierLens is a research software pilot that investigates whether satellite imagery, terrain data and machine learning can help delineate the South Lhonak Glacier in Sikkim, India across multiple dates. The project has produced a working experimental glacier-delineation workflow and visual evidence products. It is not yet an independently validated scientific study or a GLOF forecasting system.')
paragraph(doc, 'The main achievement so far is a reproducible South Lhonak pilot using three Sentinel 2 dates. Two models were trained on 2017 and 2019 imagery and tested on 2022. XGBoost currently performs better than the original Random Forest when compared with the project’s experimental derived labels. The next work is to obtain independently reviewed glacier and lake boundaries, diagnose model errors, and validate change measurements.', bold_lead='The main achievement so far is')

heading(doc, 'Research question and project objectives')
paragraph(doc, 'The long term goal is to build a transparent system for analysing glacier and lake changes using satellite imagery. The current pilot is deliberately narrower: it focuses on one glacier and one lake system before extending to additional sites.')
add_table(doc,
          ['Objective', 'Current status', 'What has been achieved'],
          [
              ['Prepare comparable multi date inputs', 'Completed for the pilot', 'Three quality accepted South Lhonak feature stacks and valid masks were prepared.'],
              ['Delineate glacier extent with machine learning', 'Experimental pilot completed', 'Random Forest and XGBoost glacier classifiers were trained and run for 2022.'],
              ['Inspect prediction errors', 'Completed for XGBoost pilot', 'Full scene shared valid footprint metrics and TP FP FN error maps were exported.'],
              ['Delineate lake extent', 'Not completed', 'Only spectral lake candidates exist; no separately evaluated lake model exists.'],
              ['Measure validated retreat and lake change', 'Not completed', 'Only provisional input boundary area changes are available.'],
              ['Multi site analysis and GLOF susceptibility', 'Not started', 'This remains outside the current experimental pilot.'],
          ], [2.1, 1.45, 3.85])

heading(doc, 'Study site and input data')
paragraph(doc, 'The pilot study site is South Lhonak Glacier and its proglacial lake in Sikkim, India. Analysis uses a fixed study area and a common 20 metre grid so that dates can be compared consistently.')
add_table(doc,
          ['Input', 'Use in the project', 'Status'],
          [
              ['Sentinel 2 surface reflectance', 'Optical satellite bands used to identify snow ice water and terrain patterns.', 'Used for 2017 2019 and 2022.'],
              ['Copernicus DEM', 'Elevation slope and aspect terrain features.', 'Used in all pilot feature stacks.'],
              ['Historical RGI glacier inventory', 'Historical spatial reference and project identity cross check.', 'Not treated as a current truth boundary.'],
              ['GEE generated lake and glacier candidates', 'Rule based candidate shapes for inspection only.', 'Not independently validated.'],
          ], [1.65, 3.7, 2.05])

heading(doc, 'Dates and experimental split')
add_table(doc,
          ['Date', 'Role', 'Reason'],
          [
              ['2017 11 19', 'Training', 'Earliest approved Sentinel 2 surface reflectance pilot date.'],
              ['2019 10 15', 'Training', 'Second late post monsoon training observation.'],
              ['2022 11 30', 'Chronological test and diagnostic date', 'Held out from initial training, then used for model comparison and full scene inspection.'],
          ], [1.25, 1.7, 4.45])
paragraph(doc, 'Because 2022 has now been used to compare the Random Forest and XGBoost models, it should no longer be described as an untouched final test date. A future date with independent labels must be kept as a new final holdout.')

heading(doc, 'Data preparation and feature engineering')
paragraph(doc, 'Each date was processed into a 13 band feature stack. Invalid pixels caused by cloud shadow or no data are kept as ignored rather than being incorrectly labelled as background. Feature order is frozen in the pipeline so that training and prediction use the same data definition.')
add_table(doc,
          ['Feature group', 'Features'],
          [
              ['Sentinel 2 reflectance', 'B2 B3 B4 B8 B11'],
              ['Spectral indices', 'NDVI NDWI MNDWI NDSI B8 B11 ratio'],
              ['Terrain', 'Elevation slope aspect'],
              ['Label values', '0 non glacier 1 glacier 255 ignored invalid or uncertain'],
          ], [1.8, 5.6])
paragraph(doc, 'The initial glacier labels were created from owner approved derived boundaries that use historical RGI geometry adjusted with dated spectral lake candidates. This makes them useful for experimental model development, but they are not independent scientific ground truth.')

heading(doc, 'Techniques used')
paragraph(doc, 'The project combines geospatial processing with pixel based machine learning. The pipeline can be summarised as: Sentinel 2 imagery and DEM data → cloud and invalid pixel masking → 13 band feature raster → experimental glacier labels → model training → 2022 probability map and glacier mask → full scene error maps and metrics.')
add_table(doc,
          ['Technique', 'How it was used'],
          [
              ['Random Forest', 'Initial binary glacier versus non glacier baseline using 250 decision trees.'],
              ['XGBoost', 'Gradient boosted tree classifier using the same temporal split and feature schema as Random Forest.'],
              ['Google Earth Engine', 'Created threshold based glacier and lake candidate polygons from Sentinel 2 and DEM data for visual inspection.'],
              ['Full scene evaluation', 'Calculates true positive false positive false negative true negative and ignored pixel maps on the common valid footprint.'],
              ['Vector comparison', 'Compares polygon overlap and projected areas instead of relying only on a single reported area.'],
          ], [1.65, 5.75])

heading(doc, 'Experimental model results')
paragraph(doc, 'The following results measure agreement with owner approved derived labels on 20,000 sampled valid pixels from the 2022 date. They do not measure independent real world accuracy.')
add_table(doc,
          ['Model', 'IoU', 'F1 Dice', 'Precision', 'Recall', 'Interpretation'],
          [
              ['Random Forest', '0.5045', '0.6706', '0.7373', '0.6150', 'Original experimental baseline.'],
              ['XGBoost', '0.5299', '0.6927', '0.7666', '0.6318', 'Modestly higher agreement on the same experimental comparison.'],
          ], [1.3, 0.7, 0.85, 0.95, 0.75, 2.0])
paragraph(doc, 'XGBoost was selected as the better experimental candidate for inspection, not as a scientifically proven final model. The observed improvement is modest and may still reflect limitations in the derived labels.')

heading(doc, 'Full scene XGBoost diagnostic')
paragraph(doc, 'A full scene evaluation was completed for the 2022 XGBoost prediction using only pixels valid in both prediction and reference masks. This prevents ignored cloud shadow and no data pixels from being counted as background.')
add_table(doc,
          ['Measure', 'Result'],
          [
              ['Shared valid pixels', '230,195'],
              ['Ignored pixels', '20,436'],
              ['True positives', '17,268'],
              ['False positives', '5,473'],
              ['False negatives', '10,314'],
              ['Full scene IoU', '0.5224'],
              ['Full scene F1 Dice', '0.6863'],
          ], [2.2, 4.3])
paragraph(doc, 'The diagnostic export includes separate true positive, false positive, false negative, true negative and ignored pixel map layers. It is useful for understanding where the model disagrees with the experimental reference, but it is not independent validation.')

heading(doc, 'Google Earth Engine candidate experiment')
paragraph(doc, 'A Google Earth Engine script was used to create rule based glacier candidate shapefiles from the 2022 Sentinel 2 scene. The rules use cloud masking, NDSI, NDVI, NDWI, elevation, a limited search area and removal of small connected components.')
add_table(doc,
          ['Candidate export', 'Area', 'Finding'],
          [
              ['Initial GEE candidate', '21.2804 km²', 'Too broad: 58 polygons captured extensive snow and bright terrain.'],
              ['Stricter GEE candidate', '12.8364 km²', 'Area is close to the experimental boundary, but spatial agreement remains low.'],
              ['Strict candidate versus XGBoost', 'IoU 0.1959', 'Similar area does not mean the same glacier geometry.'],
              ['Strict candidate versus experimental boundary', 'IoU 0.2724', 'Candidate cannot be used as a reviewed training or validation label.'],
          ], [2.0, 1.25, 3.25])
paragraph(doc, 'This experiment demonstrates an important methodological lesson: threshold based snow and ice detection can generate useful starting candidates, but it cannot reliably distinguish glacier ice from nearby seasonal snow or debris covered ice without review.')

heading(doc, 'Provisional change measurements')
paragraph(doc, 'Input boundary polygons were used to produce provisional area values. These numbers are measurements of the input geometry, not areas produced by the machine learning prediction.')
add_table(doc,
          ['Date', 'Glacier polygon area km²', 'Lake candidate area km²'],
          [
              ['2017 11 19', '12.801892', '1.149537'],
              ['2019 10 15', '12.593980', '1.358957'],
              ['2022 11 30', '12.531210', '1.713200'],
          ], [1.5, 2.4, 2.4])
paragraph(doc, 'These observations suggest changes in the input geometry over two intervals, but they must not be reported as validated glacier retreat. Terminus retreat distance, lake shoreline validation, uncertainty estimates and independent change checks are still missing.')

heading(doc, 'Current project status')
add_table(doc,
          ['Completed', 'In progress or incomplete'],
          [
              ['Three date South Lhonak input preparation and readiness gate', 'Independent glacier and lake labels'],
              ['Experimental Random Forest and XGBoost glacier models', 'Separate lake segmentation model and evaluation'],
              ['2022 prediction exports and full scene diagnostic maps', 'Uncertainty aware area and terminus change analysis'],
              ['Provisional input boundary change outputs', 'Additional verified sites and geographic holdouts'],
              ['GEE candidate shapefile experiments', 'Research release API database and evidence backed public release'],
          ], [3.25, 3.25])

heading(doc, 'Limitations and ethical scientific reporting')
bullet(doc, 'The project has an experimental working pipeline, not independently validated scientific results.')
bullet(doc, 'The current labels are derived owner approved boundaries and may share systematic errors across dates.')
bullet(doc, 'No lake model has been trained or evaluated. A glacier prediction must never be presented as a lake prediction.')
bullet(doc, 'Area similarity alone is insufficient. Maps must be checked for spatial overlap and valid coverage.')
bullet(doc, 'The project does not predict GLOF probability, flood routing, exposure or issue operational alerts.')
bullet(doc, 'Human glaciology or GIS review is required before promoting any boundary to independent reference data.')

heading(doc, 'Future steps')
add_table(doc,
          ['Priority', 'Next step', 'Why it matters'],
          [
              ['1', 'Create independently reviewed glacier and lake boundaries for the pilot dates.', 'Provides defensible validation data.'],
              ['2', 'Use the TP FP FN maps to identify debris snow shadow and alignment errors.', 'Directs improvements toward the real failure mode.'],
              ['3', 'Build a separate lake baseline and lake model.', 'Completes the glacier lake architecture.'],
              ['4', 'Tune only on a separate validation date and retain a fresh final holdout.', 'Avoids reporting over fitted test performance.'],
              ['5', 'Calculate area terminus and lake change with uncertainty and detection limits.', 'Converts candidate geometry into defensible measurements.'],
              ['6', 'Extend only to verified additional sites and use geographic holdouts.', 'Tests whether the approach generalises beyond South Lhonak.'],
          ], [0.6, 3.45, 2.45])

heading(doc, 'Conclusion')
paragraph(doc, 'GlacierLens has achieved a meaningful experimental milestone: it can prepare multi date satellite and terrain data, train and compare glacier classifiers, export prediction layers, generate reproducible error maps and explore GEE derived candidate polygons. XGBoost presently outperforms the Random Forest on the project’s derived label comparison. However, the project is not ready to claim validated glacier retreat, lake growth, regional patterns or GLOF risk. The essential next milestone is independently reviewed glacier and lake reference data.')

heading(doc, 'Key evidence and reproducibility files')
for item in [
    'Implementation plan: docs/full-architecture-implementation-plan.md',
    'XGBoost metrics: data/derived/phase2/models/segmentation/phase-2-seg-xgb-v1.0/metrics.json',
    'Full scene evaluation: data/derived/phase2/evaluations/phase-2-seg-xgb-v1.0/south-lhonak/2022-11-30/full_scene_evaluation.json',
    'Training script: pipelines/geoai/train_segmentation.py',
    'Full scene evaluator: pipelines/geoai/evaluate_segmentation_full_scene.py',
]:
    bullet(doc, item)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run('GlacierLens South Lhonak experimental pilot | Teacher review report | 27 September 2026')
footer_run.font.size = Pt(8)
footer_run.font.color.rgb = RGBColor(90, 90, 90)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
